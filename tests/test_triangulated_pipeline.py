import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import image_ocr_match_report as report


class TriangulatedPipelineTests(unittest.TestCase):
    def _ai_result(self, side: str, filename: str, language_code: str, language_name: str, objects):
        return report.OCRNormalizedResult(
            image_filename=filename,
            language_code=language_code,
            language_name=language_name,
            side=side,  # type: ignore[arg-type]
            source_engine="ai_image_ocr",
            model="gpt-4.1-mini",
            prompt_version=report.AI_IMAGE_OCR_PROMPT_VERSION,
            ui_objects=objects,
        )

    def _obj(self, oid: str, text: str, role: str = "menu_label", order: int = 1):
        return report.NormalizedGUIObject(
            object_id=oid,
            normalized_text=text,
            gui_role=role,  # type: ignore[arg-type]
            reading_order=order,
            row_group=order,
        )
    def test_ocr_candidate_package_keeps_pass_metadata(self):
        line = report.OCRLine(
            pair_id="0001",
            side="source",
            image_filename="enis01ct009b.eps",
            ocr_pass="psm6_gray",
            preprocessing="gray",
            psm="6",
            line_no=1,
            raw_line="Volume",
            cleaned_line="Volume",
            avg_tesseract_confidence=91.0,
            bbox_left=1,
            bbox_top=2,
            bbox_right=30,
            bbox_bottom=10,
            kept_or_filtered="kept",
            selected_pass=True,
            line_quality_score=3.0,
        )
        evidence = report.ClassicOCRResult([], [line], ["Volume"], "Volume", [])
        package = report.ocr_candidates_package("enis01ct009b.eps", "en", "English", "source", evidence)
        self.assertEqual(package["ocr_passes"][0]["pass_name"], "psm6_gray")
        self.assertTrue(package["ocr_passes"][0]["selected_pass"])
        self.assertEqual(package["ocr_passes"][0]["lines"][0]["cleaned_line"], "Volume")

    def test_triangulation_selects_llm_ai_agreement(self):
        classic = report.ClassicOCRResult([], [], ["Hanger6"], "Hanger6", [])
        llm = report.OCRNormalizedResult(
            image_filename="huis01ct009b.eps",
            language_code="hu",
            language_name="Hungarian",
            side="target",
            source_engine="llm_ocr_normalization",
            model="gpt-4.1-mini",
            prompt_version=report.LLM_OCR_NORMALIZATION_PROMPT_VERSION,
            ui_objects=[
                report.NormalizedGUIObject(
                    object_id="llm_001",
                    normalized_text="Hangerő",
                    normalization_confidence=0.92,
                    ocr_evidence_confidence=0.80,
                    reading_order=1,
                )
            ],
        )
        ai = report.OCRNormalizedResult(
            image_filename="huis01ct009b.eps",
            language_code="hu",
            language_name="Hungarian",
            side="target",
            source_engine="ai_image_ocr",
            model="gpt-4.1-mini",
            prompt_version=report.AI_IMAGE_OCR_PROMPT_VERSION,
            ui_objects=[
                report.NormalizedGUIObject(
                    object_id="ai_001",
                    normalized_text="Hangerő",
                    normalization_confidence=0.95,
                    ocr_evidence_confidence=0.95,
                    reading_order=1,
                    correction_type="vision_read",
                )
            ],
        )
        rows = report.triangulate_side("0001", "target", "huis01ct009b.eps", classic, llm, ai)
        self.assertEqual(rows[0].agreement_status, "llm_ai_agree")
        self.assertEqual(rows[0].selected_final_text, "Hangerő")
        self.assertEqual(rows[0].selected_source, "merged")

    def test_multi_sheet_workbook_contains_expected_sheets(self):
        pair = report.ImagePair(
            pair_id="0001",
            source_image_filename="enis01ct009b.eps",
            source_image_path="src",
            target_language_code="hu",
            target_language_name="Hungarian",
            target_image_filename="huis01ct009b.eps",
            target_image_path="trg",
            image_match_method="exact_normalized",
            image_match_confidence=1.0,
            image_match_status="OK",
        )
        pipe_report = report.PipelineReport(
            image_pairs=[pair],
            classic_lines=[],
            classic_tokens=[],
            raw_ocr_full_text=[],
            llm_results=[],
            ai_results=[],
            triangulation_rows=[],
            final_matches=[],
            summary={"ocr_engine": "classic"},
        )
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "report.xlsx"
            report.write_multi_sheet_report(pipe_report, out)
            wb = load_workbook(out)
            self.assertIn("Image pairs", wb.sheetnames)
            self.assertIn("Raw OCR - Classic Lines", wb.sheetnames)
            self.assertIn("Final Matches", wb.sheetnames)
            self.assertIn("Summary", wb.sheetnames)

    def test_lane_classification_and_raw_corroboration(self):
        obj_label = report.NormalizedGUIObject(object_id="a", normalized_text="Navigation", gui_role="menu_label")
        obj_value = report.NormalizedGUIObject(object_id="b", normalized_text="53 min • 72 km", gui_role="value")
        obj_masked = report.NormalizedGUIObject(object_id="c", normalized_text="**** ****")
        self.assertEqual(report.classify_object_lane(obj_label), "translatable_gui")
        self.assertEqual(report.classify_object_lane(obj_value), "value_only")
        self.assertEqual(report.classify_object_lane(obj_masked), "masked_text")

        evidence = report.ClassicOCRResult([], [], [], "", [], {"psm6_gray": "Navigáció Média Telefon Jármű"})
        status, _reason = report.corroborate_with_raw_ocr("Navigáció", evidence)
        self.assertIn(status, {"confirmed", "weakly_confirmed"})
        status2, _reason2 = report.corroborate_with_raw_ocr("Forgalom", evidence)
        self.assertEqual(status2, "not_confirmed")

    def test_short_warning_status_is_matchable(self):
        src_warning = report.NormalizedGUIObject(
            object_id="s1",
            normalized_text="BRAKE",
            gui_role="status_bar",
            reading_order=1,
            row_group=1,
            screen_area="bottom",
            ocr_evidence_confidence=0.95,
            normalization_confidence=0.95,
            is_translatable_gui_string=True,
        )
        trg_warning = report.NormalizedGUIObject(
            object_id="t1",
            normalized_text="F\u00c9K!",
            gui_role="status_bar",
            reading_order=1,
            row_group=1,
            screen_area="bottom",
            ocr_evidence_confidence=0.99,
            normalization_confidence=0.99,
            is_translatable_gui_string=True,
        )
        self.assertEqual(report.classify_object_lane(src_warning), "translatable_gui")
        self.assertEqual(report.classify_object_lane(trg_warning), "translatable_gui")

        source_ai = self._ai_result("source", "enis07ct033a.eps", "en", "English", [src_warning])
        target_ai = self._ai_result("target", "huis07ct033a.eps", "hu", "Hungarian", [trg_warning])
        src_classic = report.ClassicOCRResult([], [], [], "", [], {"psm11_bw": "BRAKE"})
        trg_classic = report.ClassicOCRResult([], [], [], "", [], {"psm11_bw": "Kdi F\u00c9K\nFEK"})
        out = report.ai_validated_matches(
            pair_id="0110",
            source_image="enis07ct033a.eps",
            target_image="huis07ct033a.eps",
            target_language_name="Hungarian",
            source_ai=source_ai,
            target_ai=target_ai,
            source_classic=src_classic,
            target_classic=trg_classic,
            model="gpt-4.1-mini",
            api_key=None,
            timeout_sec=1,
            max_retries=0,
        )
        self.assertEqual(len(out), 1)
        self.assertEqual(out[0].source_text, "BRAKE")
        self.assertEqual(out[0].target_text, "F\u00c9K!")
        self.assertIn(out[0].raw_ocr_validation_status, {"confirmed", "weakly_confirmed"})

    def test_no_target_reuse_in_accepted_rows(self):
        matches = [
            report.GUIMatch(
                pair_id="0001",
                source_image_filename="s.eps",
                target_image_filename="t.eps",
                source_object_id="s1",
                target_object_id="t1",
                source_text="A",
                target_text="AA",
                match_type="semantic_translation",
                semantic_confidence=0.95,
                layout_confidence=0.95,
                overall_confidence=0.95,
                status="matched",
                rationale="x",
            ),
            report.GUIMatch(
                pair_id="0001",
                source_image_filename="s.eps",
                target_image_filename="t.eps",
                source_object_id="s2",
                target_object_id="t1",
                source_text="B",
                target_text="AA",
                match_type="semantic_translation",
                semantic_confidence=0.90,
                layout_confidence=0.90,
                overall_confidence=0.80,
                status="matched",
                rationale="x",
            ),
        ]
        report.enforce_one_to_one_target_assignment(matches)
        statuses = [m.status for m in matches]
        self.assertEqual(statuses.count("matched"), 1)
        self.assertIn("uncertain", statuses)

    def test_no_accept_when_mismatch_or_not_confirmed(self):
        m1 = report.GUIMatch(
            pair_id="0001",
            source_image_filename="s.eps",
            target_image_filename="t.eps",
            source_object_id="s1",
            target_object_id="t1",
            source_text="Weather",
            target_text="Forgalom",
            match_type="semantic_translation",
            semantic_confidence=0.95,
            layout_confidence=0.9,
            overall_confidence=0.92,
            status="matched",
            rationale="x",
            semantic_check_status="mismatch",
            raw_ocr_validation_status="confirmed",
        )
        report.apply_final_quality_gates(m1)
        self.assertEqual(m1.status, "needs_review")

        m2 = report.GUIMatch(
            pair_id="0001",
            source_image_filename="s.eps",
            target_image_filename="t.eps",
            source_object_id="s2",
            target_object_id="t2",
            source_text="Message",
            target_text="Üzenet",
            match_type="semantic_translation",
            semantic_confidence=0.95,
            layout_confidence=0.9,
            overall_confidence=0.92,
            status="matched",
            rationale="x",
            semantic_check_status="ok",
            raw_ocr_validation_status="not_confirmed",
        )
        report.apply_final_quality_gates(m2)
        self.assertEqual(m2.status, "needs_review")

    def test_pair_a_values_only_suppressed(self):
        source_ai = self._ai_result(
            "source",
            "enis03ct033b.eps",
            "en",
            "English",
            [
                self._obj("s1", "1.0 km", "value", 1),
                self._obj("s2", "53 min • 72 km", "value", 2),
                self._obj("s3", "16:45", "value", 3),
            ],
        )
        target_ai = self._ai_result(
            "target",
            "huis03ct033b.eps",
            "hu",
            "Hungarian",
            [
                self._obj("t1", "1.0 km", "value", 1),
                self._obj("t2", "53 perc • 72 km", "value", 2),
                self._obj("t3", "16:45", "value", 3),
            ],
        )
        src_classic = report.ClassicOCRResult([], [], [], "", [], {"psm6_gray": "1.0 km 53 min 72 km 16:45"})
        trg_classic = report.ClassicOCRResult([], [], [], "", [], {"psm6_gray": "1.0 km 53 perc 72 km 16:45"})
        out = report.ai_validated_matches(
            pair_id="0059",
            source_image="enis03ct033b.eps",
            target_image="huis03ct033b.eps",
            target_language_name="Hungarian",
            source_ai=source_ai,
            target_ai=target_ai,
            source_classic=src_classic,
            target_classic=trg_classic,
            model="gpt-4.1-mini",
            api_key=None,
            timeout_sec=1,
            max_retries=0,
        )
        self.assertEqual(out, [])

    def test_pair_b_keeps_phone_screen_labels(self):
        labels_en = ["Favourites", "Recents", "Contacts", "Keypad", "Message", "New Message"]
        labels_hu = ["Kedvencek", "Legutóbbiak", "Névjegyek", "Billentyűzet", "Üzenet", "Új üzenet"]
        src_objs = [self._obj(f"s{i}", t, "menu_label", i) for i, t in enumerate(labels_en, start=1)]
        src_objs.append(self._obj("sv", "05/09/2025 Mobile", "value", 7))
        trg_objs = [self._obj(f"t{i}", t, "menu_label", i) for i, t in enumerate(labels_hu, start=1)]
        trg_objs.append(self._obj("tv", "05/09/2025 Mobil", "value", 7))
        source_ai = self._ai_result("source", "enis05ct016c.eps", "en", "English", src_objs)
        target_ai = self._ai_result("target", "huis05ct016b.eps", "hu", "Hungarian", trg_objs)
        src_classic = report.ClassicOCRResult([], [], [], "", [], {"psm6_gray": " ".join(labels_en)})
        trg_classic = report.ClassicOCRResult([], [], [], "", [], {"psm6_gray": " ".join(labels_hu)})
        out = report.ai_validated_matches(
            pair_id="0103",
            source_image="enis05ct016c.eps",
            target_image="huis05ct016b.eps",
            target_language_name="Hungarian",
            source_ai=source_ai,
            target_ai=target_ai,
            source_classic=src_classic,
            target_classic=trg_classic,
            model="gpt-4.1-mini",
            api_key=None,
            timeout_sec=1,
            max_retries=0,
        )
        self.assertEqual([m.source_text for m in out], labels_en)
        self.assertTrue(all("Mobile" not in m.source_text for m in out))

    def test_pair_c_weather_forgalom_blocked(self):
        source_ai = self._ai_result(
            "source",
            "enis01ct013a.eps",
            "en",
            "English",
            [self._obj("s1", "Weather", "menu_label", 1)],
        )
        target_ai = self._ai_result(
            "target",
            "huis01ct013a.eps",
            "hu",
            "Hungarian",
            [self._obj("t1", "Forgalom", "menu_label", 1)],
        )
        src_classic = report.ClassicOCRResult([], [], [], "", [], {"psm6_gray": "Navigation Media Weather Phone Vehicle"})
        trg_classic = report.ClassicOCRResult([], [], [], "", [], {"psm6_gray": "Navigáció járás Telefon Jármű"})

        original_call = report.call_openai_structured
        original_sem = report.semantic_quality_check
        try:
            def fake_call(*args, **kwargs):
                return report.GUIObjectMatchResponse(
                    pair_id="0012",
                    source_image="enis01ct013a.eps",
                    target_image="huis01ct013a.eps",
                    matches=[
                        report.GUIMatch(
                            pair_id="0012",
                            source_image_filename="enis01ct013a.eps",
                            target_image_filename="huis01ct013a.eps",
                            source_object_id="s1",
                            target_object_id="t1",
                            source_text="Weather",
                            target_text="Forgalom",
                            match_type="semantic_translation",
                            semantic_confidence=0.8,
                            layout_confidence=0.8,
                            overall_confidence=0.8,
                            status="matched",
                            rationale="test",
                        )
                    ],
                )

            def fake_sem(**kwargs):
                return report.SemanticPairCheck(
                    semantic_check_status="mismatch",
                    semantic_check_confidence=0.98,
                    semantic_check_reason="Meaning differs significantly.",
                )

            report.call_openai_structured = fake_call  # type: ignore[assignment]
            report.semantic_quality_check = fake_sem  # type: ignore[assignment]

            out = report.ai_validated_matches(
                pair_id="0012",
                source_image="enis01ct013a.eps",
                target_image="huis01ct013a.eps",
                target_language_name="Hungarian",
                source_ai=source_ai,
                target_ai=target_ai,
                source_classic=src_classic,
                target_classic=trg_classic,
                model="gpt-4.1-mini",
                api_key="dummy",
                timeout_sec=1,
                max_retries=0,
            )
        finally:
            report.call_openai_structured = original_call  # type: ignore[assignment]
            report.semantic_quality_check = original_sem  # type: ignore[assignment]

        self.assertEqual(len(out), 1)
        self.assertEqual(out[0].status, "needs_review")
        self.assertEqual(out[0].semantic_check_status, "mismatch")
        self.assertEqual(out[0].raw_ocr_validation_status, "not_confirmed")


if __name__ == "__main__":
    unittest.main()

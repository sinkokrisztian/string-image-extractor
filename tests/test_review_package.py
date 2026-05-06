from pathlib import Path
import tempfile
import unittest

import pandas as pd
from openpyxl import load_workbook
from PIL import Image

from scripts.build_review_package import build_for_workbook


def _write_png(path: Path, color: tuple[int, int, int]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (12, 8), color).save(path)


class ReviewPackageTests(unittest.TestCase):
    def test_build_review_package_backfills_mm_strings_and_links_pngs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            src_png = tmp_path / "images" / "source" / "enis01ct001a.png"
            tgt_png = tmp_path / "images" / "target" / "csis01ct001a.png"
            _write_png(src_png, (220, 20, 20))
            _write_png(tgt_png, (20, 80, 220))

            mm_source = tmp_path / "MM_strings_EN.xlsx"
            pd.DataFrame({"English": ["Save", "Cancel"]}).to_excel(mm_source, index=False)

            report = tmp_path / "report_gui.xlsx"
            with pd.ExcelWriter(report, engine="openpyxl") as writer:
                pd.DataFrame().to_excel(writer, sheet_name="MM strings", index=False)
                pd.DataFrame(
                    [
                        {
                            "source_text": "Save",
                            "target_text": "Ulozit",
                            "status": "matched",
                            "overall_confidence": 0.98,
                            "semantic_confidence": 0.98,
                            "rationale": "test match",
                            "review_reason": "",
                            "semantic_check_reason": "",
                            "source_image_path": str(src_png),
                            "target_image_path": str(tgt_png),
                        }
                    ]
                ).to_excel(writer, sheet_name="Final Matches", index=False)
                pd.DataFrame([{"target_language_code": "cs"}]).to_excel(writer, sheet_name="Image pairs", index=False)

            out = build_for_workbook(
                report,
                tmp_path / "to_review",
                gs_cmd="unused",
                dpi=72,
                mm_strings_xlsx=mm_source,
            )

            wb = load_workbook(out)
            ws = wb["MM strings"]
            headers = [cell.value for cell in ws[1]]
            status_col = headers.index("status") + 1
            src_col = headers.index("source_image") + 1
            tgt_col = headers.index("target_image") + 1

            self.assertEqual(out.name, "MM_strings_review_ENCS.xlsx")
            self.assertEqual(ws.max_row, 3)
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertEqual(ws.auto_filter.ref, "A1:I3")
            self.assertEqual(ws.cell(row=2, column=status_col).value, "matched")
            self.assertEqual(ws.cell(row=3, column=status_col).value, "unmatched")
            self.assertIsNotNone(ws.cell(row=2, column=src_col).hyperlink)
            self.assertIsNotNone(ws.cell(row=2, column=tgt_col).hyperlink)
            self.assertTrue((tmp_path / "to_review" / "cs" / "source" / "enis01ct001a.png").exists())
            self.assertTrue((tmp_path / "to_review" / "cs" / "target" / "csis01ct001a.png").exists())

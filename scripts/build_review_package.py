from __future__ import annotations

import argparse
import os
import re
import subprocess
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image


def normalize_key(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "")).strip().lower()


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_target_lang_from_filename(path: Path) -> str:
    m = re.search(r"EN-?([A-Za-z]{2})", path.stem, flags=re.IGNORECASE)
    if m:
        return m.group(1).lower()
    m = re.search(r"EN([A-Za-z]{2})", path.stem, flags=re.IGNORECASE)
    if m:
        return m.group(1).lower()
    return "xx"


def infer_target_lang(xlsx: Path, explicit_lang: Optional[str] = None) -> str:
    if explicit_lang:
        lang = explicit_lang.strip().lower()
        if lang:
            return lang

    filename_lang = parse_target_lang_from_filename(xlsx)
    if filename_lang != "xx":
        return filename_lang

    try:
        image_pairs = pd.read_excel(xlsx, sheet_name="Image pairs")
        if "target_language_code" in image_pairs.columns:
            for value in image_pairs["target_language_code"].tolist():
                lang = cell_text(value).lower()
                if lang:
                    return lang
    except Exception:
        pass
    return "xx"


def ensure_png(src: Path, dst: Path, gs_cmd: str, dpi: int) -> None:
    dst.parent.mkdir(parents=True, exist_ok=True)
    if dst.exists():
        crop_png_to_content(dst)
        return
    ext = src.suffix.lower()
    if ext in {".png", ".jpg", ".jpeg", ".bmp"}:
        dst.write_bytes(src.read_bytes())
        crop_png_to_content(dst)
        return
    if ext == ".eps":
        cmd = [
            gs_cmd,
            "-dSAFER",
            "-dBATCH",
            "-dNOPAUSE",
            "-sDEVICE=pngalpha",
            f"-r{dpi}",
            f"-sOutputFile={str(dst)}",
            str(src),
        ]
        subprocess.run(cmd, check=True, capture_output=True)
        crop_png_to_content(dst)
        return
    raise RuntimeError(f"Unsupported image extension for review export: {src}")


def crop_png_to_content(path: Path, pad: int = 8) -> None:
    # Crop transparent/empty margins around rendered screenshots.
    with Image.open(path) as img:
        rgba = img.convert("RGBA")
        alpha = rgba.split()[3]
        bbox = alpha.getbbox()
        if not bbox:
            return
        left, top, right, bottom = bbox
        # Fallback for fully opaque images with white canvas: detect non-white area.
        if left == 0 and top == 0 and right == rgba.width and bottom == rgba.height:
            rgb = rgba.convert("RGB")
            bg = rgb.getpixel((0, 0))
            tol = 12
            px = rgb.load()
            min_x, min_y = rgb.width, rgb.height
            max_x, max_y = -1, -1
            for y in range(rgb.height):
                for x in range(rgb.width):
                    r, g, b = px[x, y]
                    if abs(r - bg[0]) > tol or abs(g - bg[1]) > tol or abs(b - bg[2]) > tol:
                        if x < min_x:
                            min_x = x
                        if y < min_y:
                            min_y = y
                        if x > max_x:
                            max_x = x
                        if y > max_y:
                            max_y = y
            if max_x >= min_x and max_y >= min_y:
                left, top, right, bottom = min_x, min_y, max_x + 1, max_y + 1
        left = max(0, left - pad)
        top = max(0, top - pad)
        right = min(rgba.width, right + pad)
        bottom = min(rgba.height, bottom + pad)
        cropped = rgba.crop((left, top, right, bottom))
        cropped.save(path)


def choose_best_row(group: pd.DataFrame) -> Optional[pd.Series]:
    if group.empty:
        return None
    rank = {"matched": 4, "needs_review": 3, "uncertain": 2, "unmatched": 1}
    tmp = group.copy()
    tmp["_rank"] = tmp["status"].map(lambda s: rank.get(str(s), 0))
    tmp["_conf"] = pd.to_numeric(tmp.get("overall_confidence"), errors="coerce").fillna(0.0)
    tmp = tmp.sort_values(["_rank", "_conf"], ascending=[False, False])
    return tmp.iloc[0]


def read_manual_english_strings(mm_xlsx: Path) -> List[str]:
    strings: List[str] = []
    if not mm_xlsx.exists():
        return strings
    with pd.ExcelFile(mm_xlsx) as xls:
        for sheet in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet)
            if df.empty:
                continue
            cols = [str(c).strip() for c in df.columns]
            lower_cols = [c.lower() for c in cols]
            en_col_idx = -1
            for idx, col in enumerate(lower_cols):
                if col in {"en", "english", "source", "source_en", "om_en"}:
                    en_col_idx = idx
                    break
            if en_col_idx < 0:
                for idx, col in enumerate(lower_cols):
                    if "en" in col or "english" in col:
                        en_col_idx = idx
                        break
            if en_col_idx < 0:
                en_col_idx = 0
            for value in df[cols[en_col_idx]].tolist():
                text = cell_text(value)
                if text:
                    strings.append(text)

    deduped: List[str] = []
    seen = set()
    for text in strings:
        key = normalize_key(text)
        if not key or key in seen:
            continue
        seen.add(key)
        deduped.append(text)
    return deduped


def merge_unique_text(rows: List[pd.Series], field: str) -> str:
    values: List[str] = []
    seen = set()
    for row in rows:
        text = cell_text(row.get(field, ""))
        key = normalize_key(text)
        if not key or key in seen:
            continue
        seen.add(key)
        values.append(text)
    return " | ".join(values)


def build_manual_rows_from_matches(fm: pd.DataFrame, mm_xlsx: Path) -> List[Dict[str, Any]]:
    manual_strings = read_manual_english_strings(mm_xlsx)
    if not manual_strings:
        return []

    fm = fm.copy()
    if "source_text" not in fm.columns:
        return []
    fm["source_key"] = fm["source_text"].map(cell_text).map(normalize_key)
    by_key: Dict[str, pd.DataFrame] = {k: g for k, g in fm.groupby("source_key") if k}

    rows: List[Dict[str, Any]] = []
    for source in manual_strings:
        key = normalize_key(source)
        group = by_key.get(key)
        hits = [] if group is None else [row for _, row in group.iterrows()]
        targets: List[str] = []
        for row in hits:
            target = cell_text(row.get("target_text", ""))
            if target and target not in targets:
                targets.append(target)

        status = "unmatched"
        rationale = ""
        review_reason = ""
        semantic_check_reason = ""
        if hits:
            if len(hits) == 1:
                best = hits[0]
                status = cell_text(best.get("status", "")) or "unmatched"
                rationale = cell_text(best.get("rationale", ""))
                review_reason = cell_text(best.get("review_reason", ""))
                semantic_check_reason = cell_text(best.get("semantic_check_reason", ""))
            else:
                status = "uncertain"
                rationale = merge_unique_text(hits, "rationale")
                review_reason = merge_unique_text(hits, "review_reason")
                semantic_check_reason = merge_unique_text(hits, "semantic_check_reason")

        rows.append(
            {
                "om_en_string": source,
                "target_equivalent_candidates": " | ".join(targets),
                "status": status,
                "rationale": rationale,
                "review_reason": review_reason,
                "semantic_check_reason": semantic_check_reason,
                "match_count": len(hits),
            }
        )
    return rows


def format_mm_strings_sheet(ws: Any) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "om_en_string": 34,
        "target_equivalent_candidates": 48,
        "status": 14,
        "rationale": 54,
        "review_reason": 46,
        "semantic_check_reason": 58,
        "match_count": 12,
        "source_image": 18,
        "target_image": 18,
    }
    wrap_headers = {
        "om_en_string",
        "target_equivalent_candidates",
        "rationale",
        "review_reason",
        "semantic_check_reason",
    }

    header = {str(ws.cell(row=1, column=c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    for name, col_idx in header.items():
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = widths.get(name, max(12, min(60, len(name) + 2)))
        if name in wrap_headers:
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    status_col = header.get("status")
    if not status_col or ws.max_row < 2:
        return

    # Clear old conditional-formatting rules so regenerated sheets stay tidy.
    try:
        ws.conditional_formatting._cf_rules.clear()
    except Exception:
        pass

    fills = {
        "matched": PatternFill(start_color="FFD9EAD3", end_color="FFD9EAD3", fill_type="solid"),
        "needs_review": PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid"),
        "uncertain": PatternFill(start_color="FFFCE5CD", end_color="FFFCE5CD", fill_type="solid"),
        "unmatched": PatternFill(start_color="FFF4CCCC", end_color="FFF4CCCC", fill_type="solid"),
    }
    status_letter = get_column_letter(status_col)
    max_col = get_column_letter(ws.max_column)
    target_range = f"A2:{max_col}{ws.max_row}"
    for status, fill in fills.items():
        ws.conditional_formatting.add(
            target_range,
            FormulaRule(formula=[f'${status_letter}2="{status}"'], stopIfTrue=False, fill=fill),
        )


def build_for_workbook(
    xlsx: Path,
    review_root: Path,
    gs_cmd: str,
    dpi: int,
    mm_strings_xlsx: Path,
    target_lang: Optional[str] = None,
) -> Path:
    lang = infer_target_lang(xlsx, target_lang)
    lang_dir = review_root / lang
    src_dir = lang_dir / "source"
    tgt_dir = lang_dir / "target"
    out_xlsx = review_root / f"MM_strings_review_EN{lang.upper()}.xlsx"
    review_root.mkdir(parents=True, exist_ok=True)

    try:
        mm = pd.read_excel(xlsx, sheet_name="MM strings")
    except ValueError:
        mm = pd.DataFrame()
    fm = pd.read_excel(xlsx, sheet_name="Final Matches")
    mm_headers = {str(c).strip() for c in mm.columns}
    generated_mm_rows: List[Dict[str, Any]] = []
    if mm.empty or "om_en_string" not in mm_headers:
        generated_mm_rows = build_manual_rows_from_matches(fm, mm_strings_xlsx)
        if generated_mm_rows:
            mm = pd.DataFrame(generated_mm_rows)

    # Build source_text -> candidate rows for reviewable statuses.
    reviewable = fm[fm["status"].astype(str).isin({"matched", "uncertain", "needs_review"})].copy()
    reviewable["source_key"] = reviewable["source_text"].astype(str).map(normalize_key)
    by_key: Dict[str, pd.DataFrame] = {k: g for k, g in reviewable.groupby("source_key")}

    # Copy workbook first to preserve styles.
    wb = load_workbook(xlsx)
    wb.save(out_xlsx)
    wb = load_workbook(out_xlsx)
    if generated_mm_rows:
        if "MM strings" in wb.sheetnames:
            del wb["MM strings"]
        ws = wb.create_sheet("MM strings", 0)
        for col_idx, name in enumerate(mm.columns, start=1):
            ws.cell(row=1, column=col_idx).value = name
        for row_idx, row in enumerate(mm.to_dict("records"), start=2):
            for col_idx, name in enumerate(mm.columns, start=1):
                ws.cell(row=row_idx, column=col_idx).value = row.get(name, "")
    else:
        ws = wb["MM strings"]
    header = {str(ws.cell(row=1, column=c).value or "").strip(): c for c in range(1, ws.max_column + 1)}

    # Add columns if missing.
    add_cols = ["source_image", "target_image"]
    next_col = ws.max_column + 1
    for col in add_cols:
        if col not in header:
            ws.cell(row=1, column=next_col).value = col
            header[col] = next_col
            next_col += 1

    for row_idx in range(2, ws.max_row + 1):
        s = str(ws.cell(row=row_idx, column=header["om_en_string"]).value or "")
        status = str(ws.cell(row=row_idx, column=header["status"]).value or "")
        if status not in {"matched", "uncertain", "needs_review"}:
            continue

        key = normalize_key(s)
        grp = by_key.get(key)
        if grp is None or grp.empty:
            continue
        best = choose_best_row(grp)
        if best is None:
            continue

        src_path = Path(str(best.get("source_image_path") or "")).expanduser()
        tgt_path = Path(str(best.get("target_image_path") or "")).expanduser()
        if not src_path.exists() or not tgt_path.exists():
            continue

        src_png = src_dir / (src_path.stem + ".png")
        tgt_png = tgt_dir / (tgt_path.stem + ".png")
        ensure_png(src_path, src_png, gs_cmd, dpi)
        ensure_png(tgt_path, tgt_png, gs_cmd, dpi)

        src_cell = ws.cell(row=row_idx, column=header["source_image"])
        tgt_cell = ws.cell(row=row_idx, column=header["target_image"])
        src_cell.value = src_png.name
        tgt_cell.value = tgt_png.name
        src_rel = Path(os.path.relpath(src_png, out_xlsx.parent)).as_posix()
        tgt_rel = Path(os.path.relpath(tgt_png, out_xlsx.parent)).as_posix()
        src_cell.hyperlink = src_rel
        tgt_cell.hyperlink = tgt_rel
        src_cell.style = "Hyperlink"
        tgt_cell.style = "Hyperlink"

    format_mm_strings_sheet(ws)
    wb.save(out_xlsx)
    return out_xlsx


def main() -> None:
    parser = argparse.ArgumentParser(description="Build reviewer package with MM strings + linked source/target PNGs.")
    parser.add_argument("--workbooks", nargs="+", required=True, help="Input workbook paths.")
    parser.add_argument("--review-root", type=Path, default=Path("to_review"), help="Output review root folder.")
    parser.add_argument("--mm-strings-xlsx", type=Path, default=Path("input") / "MM_strings_EN.xlsx")
    parser.add_argument("--ghostscript-cmd", default=r"C:\Program Files\gs\gs10.07.0\bin\gswin64c.EXE")
    parser.add_argument("--dpi", type=int, default=300)
    parser.add_argument("--target-lang", default=None, help="Optional target language code override, e.g. hu/sk/sl.")
    args = parser.parse_args()

    args.review_root.mkdir(parents=True, exist_ok=True)
    for wb in args.workbooks:
        out = build_for_workbook(
            Path(wb),
            args.review_root,
            args.ghostscript_cmd,
            args.dpi,
            args.mm_strings_xlsx,
            args.target_lang,
        )
        print(f"Built: {out}")


if __name__ == "__main__":
    main()
